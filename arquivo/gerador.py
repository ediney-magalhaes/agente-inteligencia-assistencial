from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
import tempfile
import os

import processador
import visualizador
import agente_ia
import utils


def _inserir_grafico(document, fig, largura=6):
    """Salva figura matplotlib em arquivo temporário e insere no documento."""
    with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp:
        fig.savefig(tmp.name, bbox_inches='tight', dpi=150)
        tmp_path = tmp.name
    document.add_picture(tmp_path, width=Inches(largura))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    os.remove(tmp_path)


def _inserir_imagem_epimed(document, imagem_path, largura=6):
    """Insere imagem estática fornecida pelo usuário (capturas do EPIMED)."""
    if imagem_path and os.path.exists(imagem_path):
        document.add_picture(imagem_path, width=Inches(largura))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        document.add_paragraph('[Imagem EPIMED não fornecida]')


def _fonte(document, texto='Fonte: EPIMED'):
    """Adiciona linha de fonte abaixo do gráfico."""
    p = document.add_paragraph(texto)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)


def _secao(document, titulo, nivel=1):
    """Adiciona título de seção."""
    document.add_heading(titulo, level=nivel)


def _analise(document, texto):
    """Adiciona parágrafo de análise gerado pela IA."""
    if texto:
        document.add_paragraph(texto)


def gerar_relatorio(trimestre,
                    ano,
                    hospital,
                    arquivo_notificacoes,
                    arquivo_indicadores,
                    imagens,
                    texto_obitos,
                    texto_prontuarios,
                    arquivo_psp,
                    acoes_taticas):

    document = Document()

    # =========================================================================
    # BLOCO 1 — PREPARAÇÃO DOS DADOS
    # =========================================================================

    df_notificacoes = arquivo_notificacoes if isinstance(arquivo_notificacoes, pd.DataFrame) else pd.read_excel(arquivo_notificacoes)
    df_indicadores = arquivo_indicadores if isinstance(arquivo_indicadores, pd.DataFrame) else pd.read_excel(arquivo_indicadores)
    df_corrigido = processador.carregar_validar(df_notificacoes)

    # Processamento por bloco
    (total_atual, total_tri_anterior, total_ano_anterior,
     qtde_mensal, qtde_trimestral, var_tri, var_ano,
     trimestre_atual, ano_atual, trimestre_anterior, ano_anterior,
     df_atual, df_tri_anterior) = processador.preparar_bloco1(df_corrigido)

    tabela_bloco2, contexto_ia_bloco2 = processador.preparar_bloco2(df_atual, df_tri_anterior)

    tabela_bloco3, dicionario_bloco3, contexto_ia_bloco3 = processador.preparar_bloco3(df_atual, df_tri_anterior)

    (tabela_bloco4, tabela_top3_bloco4,
     status_eventos_bloco4, contexto_ia_bloco4) = processador.preparar_bloco4(df_atual, df_tri_anterior)

    tab_atual_b5, tab_anterior_b5, ctx_ia_b5 = processador.preparar_bloco_setores(
        df_atual, df_tri_anterior, utils.COL_SETOR_NOTIFICANTE)

    tab_atual_b6, tab_anterior_b6, ctx_ia_b6 = processador.preparar_bloco_setores(
        df_atual, df_tri_anterior, utils.COL_SETOR)

    df_ano_atual_queda, df_ano_anterior_queda, media_tri_atual_queda, media_tri_anterior_queda, (
        grau_dano_queda, local_queda, tipo_queda) = processador.preparar_bloco7(
        df_atual, df_indicadores, utils.INDICADOR_QUEDA, trimestre, ano)

    df_ano_atual_lpp, df_ano_anterior_lpp, df_media_tri_atual_lpp, df_media_tri_anterior_lpp, df_ctx_lpp = processador.preparar_bloco7(
        df_atual, df_indicadores, utils.INDICADOR_LPP, trimestre, ano)

    df_ano_atual_fle, df_ano_anterior_fle, df_media_tri_atual_fle, df_media_tri_anterior_fle, df_ctx_fle = processador.preparar_bloco7(
        df_atual, df_indicadores, utils.INDICADOR_FLEBITE, trimestre, ano)

    df_ano_atual_med, df_ano_anterior_med, df_media_tri_atual_med, df_media_tri_anterior_med, df_ctx_med = processador.preparar_bloco7(
        df_atual, df_indicadores, utils.INDICADOR_MEDICACAO, trimestre, ano)

    (tabela_tri_atual_b8, tabela_tri_anterior_b8,
     tabela_mensal_b8, tabela_mensal_anterior_b8,
     df_ctx_b8) = processador.preparar_bloco8(df_corrigido)

    # =========================================================================
    # BLOCO 2 — GERAÇÃO DOS TEXTOS DE ANÁLISE (agente_ia.py)
    # =========================================================================

    texto_b1 = agente_ia.analisar_bloco1(total_atual, total_tri_anterior, var_tri, qtde_mensal, qtde_trimestral)

    texto_b2 = agente_ia.analisar_bloco2(tabela_bloco2, contexto_ia_bloco2)

    texto_b3 = agente_ia.analisar_bloco3(tabela_bloco3, contexto_ia_bloco3)

    texto_b4 = agente_ia.analisar_bloco4(tabela_bloco4, tabela_top3_bloco4, status_eventos_bloco4, contexto_ia_bloco4)

    texto_b5 = agente_ia.analisar_bloco_setores(tab_atual_b5, tab_anterior_b5, ctx_ia_b5, utils.COL_SETOR_NOTIFICANTE)

    texto_b6 = agente_ia.analisar_bloco_setores(tab_atual_b6, tab_anterior_b6, ctx_ia_b6, utils.COL_SETOR)

    texto_b7_geral = agente_ia.analisar_bloco7_geral(medias)

    texto_que = agente_ia.analisar_bloco7_queda(grau_dano_queda, local_queda, tipo_queda, df_ctx_queda, media_tri_atual_queda, media_tri_anterior_queda)

    texto_lpp = agente_ia.analisar_bloco7_lpp(df_media_tri_atual_lpp, df_media_tri_anterior_lpp, df_ano_atual_lpp, df_contexto_ia_lpp)

    texto_fle = agente_ia.analisar_bloco7_flebite(df_media_tri_atual_fle, df_media_tri_anterior_fle, df_ano_atual_fle, df_contexto_ia_fle)

    texto_med = agente_ia.analisar_bloco7_medicacao(df_media_tri_atual_med, df_media_tri_anterior_med, df_ano_atual_med, df_contexto_ia_med)

    medias = {
        'Queda':            (media_tri_atual_queda,    media_tri_anterior_queda),
        'Lesão de Pele':    (df_media_tri_atual_lpp,   df_media_tri_anterior_lpp),
        'Flebite':          (df_media_tri_atual_fle,   df_media_tri_anterior_fle),
        'Erro de medicação':(df_media_tri_atual_med,   df_media_tri_anterior_med),
    }

    texto_b8 = agente_ia.analisar_bloco8(tabela_tri_atual_b8, tabela_tri_anterior_b8, tabela_mensal_b8, tabela_mensal_anterior_b8, df_ctx_b8)

    texto_interrelacao = agente_ia.analisar_interrelacao(total_atual, trimestre_atual, ano_atual, df_atual)

    # =========================================================================
    # BLOCO 3 — MONTAGEM DO DOCUMENTO (ordem do Modelo Relatório)
    # =========================================================================

    # --- CAPA ---
    titulo = document.add_heading('RELATÓRIO TRIMESTRAL DE SEGURANÇA DO PACIENTE', level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f'{trimestre_atual}º Trimestre de {ano_atual} — {hospital}').alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    # --- OBJETIVO ---
    _secao(document, '1. OBJETIVO')
    document.add_paragraph(
        f'Este documento tem como objetivo identificar e analisar os pontos críticos '
        f'relacionados à Segurança do Paciente do {trimestre_atual}º Trimestre de {ano_atual}, '
        f'elaborar estratégias de intervenção e reportá-las à apreciação da alta gestão do {hospital}.'
    )

    # --- DEFINIÇÕES ---
    _secao(document, '2. DEFINIÇÕES')
    document.add_paragraph(
        'As notificações foram classificadas nas seguintes categorias: Circunstâncias de Risco, '
        'Near Miss, Evento Adverso, Evento Sentinela, Outra Natureza, Queixa Técnica e Segurança do Trabalho.'
    )

    tabela_def = document.add_table(rows=1, cols=2)
    tabela_def.style = 'Table Grid'
    hdr = tabela_def.rows[0].cells
    hdr[0].text = 'Classificação'
    hdr[1].text = 'Definição'
    definicoes = [
        ('Circunstância de risco', 'Condição ou situação que tem o potencial de causar danos ao paciente se não for gerenciada adequadamente.'),
        ('Near miss', 'Falha que ocorreu e que poderia ter causado dano ao paciente, mas foi identificada antes que acontecesse.'),
        ('Evento adverso', 'Incidente que atinge o paciente, categorizado por gravidade: sem dano, leve, moderado, grave ou óbito.'),
        ('Evento sentinela', 'Evento evitável, incapacidade ou morte inesperada. Requer atenção imediata e investigação.'),
        ('Queixa técnica', 'Desvios de qualidade em materiais e/ou equipamentos.'),
        ('Segurança do Trabalho', 'Relacionadas à segurança ocupacional: desvios de comportamento, práticas inseguras e acidentes.'),
        ('Outra natureza', 'Não se enquadram na segurança do paciente. Relatos pessoais ou admissões de ocorrências.'),
    ]
    for classificacao, definicao in definicoes:
        row = tabela_def.add_row().cells
        row[0].text = classificacao
        row[1].text = definicao

    # --- GRÁFICO 1 — Volume de Notificações ---
    _secao(document, f'Gráfico 1: Quantidade de Notificações NQSP — {hospital}. Brasília, {ano_atual}.')
    fig1 = visualizador.grafico_volume_notificacoes(qtde_mensal, qtde_trimestral, trimestre_atual, ano_atual)
    _inserir_grafico(document, fig1)
    _fonte(document)
    _analise(document, texto_b1)

    # --- GRÁFICO 2 — Turno (imagem EPIMED) ---
    _secao(document, f'Gráfico 2: Turno de maior número de notificações. Brasília, {ano_atual}.')
    _inserir_imagem_epimed(document, imagens.get('turno'))
    _fonte(document)
    _analise(document, texto_b2)

    # --- GRÁFICO 3 — Classificações ---
    _secao(document, f'Gráfico 3: Notificações de acordo com a classificação, {trimestre_atual}º Trimestre. Brasília, {ano_atual}.')
    fig3 = visualizador.grafico_classificacoes(tabela_bloco3, trimestre_atual, ano_atual)
    _inserir_grafico(document, fig3)
    _fonte(document)
    _analise(document, texto_b3)

    # Subgráficos por classificação
    classificacoes = [
        ('Circunstância de Risco',     'Incidente (circunstância de risco ou condições inseguras)'),
        ('Near Miss',                   'Near miss'),
        ('Incidente Sem Dano',          'Incidente sem dano'),
        ('Evento Sentinela',            'Never event/Evento sentinela (ANVISA/JCI)'),
        ('Queixas Técnicas',            'Queixa técnica'),
        ('Segurança Ocupacional',       'Segurança do Trabalho'),
        ('Outra Natureza',              'Outra natureza'),
    ]
    for titulo_sub, chave_dicionario in classificacoes:
        _secao(document, titulo_sub, nivel=2)
        fig_sub = visualizador.gerar_grafico_top3(dicionario_bloco3, chave_dicionario)
        _inserir_grafico(document, fig_sub)
        _fonte(document)

    # --- GRÁFICO 4 — Eventos Adversos ---
    _secao(document, f'Gráfico 4: Eventos adversos notificados. Brasília, {ano_atual}.')

    _secao(document, '4.1: Eventos adversos do trimestre', nivel=2)
    fig41 = visualizador.gerar_grafico_bloco4_1(tabela_top3_bloco4)
    _inserir_grafico(document, fig41)
    _fonte(document)

    _secao(document, '4.2: Evento adverso por grau do dano', nivel=2)
    _inserir_imagem_epimed(document, imagens.get('grau_dano'))
    _fonte(document)
    _analise(document, texto_b4)

    # --- GRÁFICO 5 — Setores Notificantes ---
    _secao(document, f'Gráfico 5: Setores notificantes do {trimestre_atual}º Trimestre. Brasília, {ano_atual}.')
    _inserir_imagem_epimed(document, imagens.get('setores_notificantes'))
    _fonte(document)
    _analise(document, texto_b5)

    # --- GRÁFICO 6 — Setores Notificados ---
    _secao(document, f'Gráfico 6: Setores notificados (local de ocorrência) do {trimestre_atual}º Trimestre. Brasília, {ano_atual}.')
    _inserir_imagem_epimed(document, imagens.get('setores_notificados'))
    _fonte(document)
    _analise(document, texto_b6)

    # --- GRÁFICO 7 — Indicadores de Qualidade ---
    _secao(document, f'Gráfico 7: Índices de flebite, lesão por pressão, queda e erros de medicação. Brasília, {ano_atual}.')
    
    fig7 = visualizador.gerar_grafico_bloco7_geral(medias, None)
    _inserir_grafico(document, fig7)
    _fonte(document)

    _secao(document, '1. Erro de Medicação', nivel=2)
    fig_med = visualizador.gerar_grafico_bloco7(df_ano_atual_med, df_ano_anterior_med, utils.INDICADOR_MEDICACAO, None, ano_atual)
    _inserir_grafico(document, fig_med)
    _fonte(document)
    _analise(document, texto_med)

    _secao(document, '2. Lesão por Pressão', nivel=2)
    fig_lpp = visualizador.gerar_grafico_bloco7(df_ano_atual_lpp, df_ano_anterior_lpp, utils.INDICADOR_LPP, None, ano_atual)
    _inserir_grafico(document, fig_lpp)
    _fonte(document)
    _analise(document, texto_lpp)

    _secao(document, '3. Flebite', nivel=2)
    fig_fle = visualizador.gerar_grafico_bloco7(df_ano_atual_fle, df_ano_anterior_fle, utils.INDICADOR_FLEBITE, None, ano_atual)
    _inserir_grafico(document, fig_fle)
    _fonte(document)
    _analise(document, texto_fle)

    _secao(document, '4. Queda', nivel=2)
    fig_que = visualizador.gerar_grafico_bloco7(df_ano_atual_queda, df_ano_anterior_queda, utils.INDICADOR_QUEDA, None, ano_atual)
    _inserir_grafico(document, fig_que)
    _fonte(document)
    _analise(document, texto_que)

    # --- GRÁFICO 8 — Cumprimento das Análises ---
    _secao(document, f'Gráfico 8: Cumprimento de análise de notificações. Brasília, {ano_atual}.')
    fig8 = visualizador.gerar_grafico_bloco8(
        tabela_tri_atual_b8, tabela_tri_anterior_b8,
        tabela_mensal_b8, tabela_mensal_anterior_b8, ano_atual)
    _inserir_grafico(document, fig8)
    _fonte(document)
    _analise(document, texto_b8)

    # --- SEÇÃO 4 — AUDITORIA DE ROPs ---
    _secao(document, '4. AUDITORIA DE ROPs')
    _inserir_imagem_epimed(document, imagens.get('rops'))
    document.add_paragraph(imagens.get('texto_rops', '[Análise de ROPs não fornecida]'))

    # --- SEÇÃO 5 — PROTOCOLOS GERENCIADOS ---
    _secao(document, '5. PROTOCOLOS GERENCIADOS')
    document.add_paragraph('Informar apenas se houve evento adverso relacionado a algum dos protocolos abaixo.')
    for protocolo in ['5.1 AVC', '5.2 Dor Torácica', '5.3 Sepse Adulto', '5.4 TEV']:
        _secao(document, protocolo, nivel=2)
        document.add_paragraph('[Preencher se houver ocorrência]')

    # --- SEÇÃO 6 — COMISSÃO DE ÓBITOS ---
    _secao(document, '6. COMISSÃO DE ÓBITOS')
    _analise(document, texto_obitos)

    # --- SEÇÃO 7 — COMISSÃO DE PRONTUÁRIOS ---
    _secao(document, '7. COMISSÃO DE PRONTUÁRIOS')
    _analise(document, texto_prontuarios)

    # --- SEÇÃO 9 — PLANO DE SEGURANÇA DO PACIENTE ---
    _secao(document, '9. PLANO DE SEGURANÇA DO PACIENTE — ACOMPANHAMENTO DOS PLANOS DE AÇÃO')
    if arquivo_psp is not None:
        df_psp = pd.read_excel(arquivo_psp)
        tabela_psp = document.add_table(rows=1, cols=len(df_psp.columns))
        tabela_psp.style = 'Table Grid'
        for i, col in enumerate(df_psp.columns):
            tabela_psp.rows[0].cells[i].text = str(col)
        for _, linha in df_psp.iterrows():
            row = tabela_psp.add_row().cells
            for i, val in enumerate(linha):
                row[i].text = str(val)
    else:
        document.add_paragraph('[Arquivo PSP não fornecido]')

    # --- SEÇÃO 10 — INTERRELAÇÃO E MAPEAMENTO DE RISCO ---
    _secao(document, '10. INTERRELAÇÃO ENTRE AS NOTIFICAÇÕES E MAPEAMENTO DE RISCO DA INSTITUIÇÃO')
    _analise(document, texto_interrelacao)

    document.add_paragraph('Prioridades críticas identificadas no trimestre:').runs[0].bold = True
    document.add_paragraph('[A IA identificou as prioridades no texto acima. Complemente com o mapeamento NPR abaixo.]')

    # Estrutura do mapeamento de risco — preenchida pelo profissional
    campos_npr = ['Risco', 'Fatores contribuintes', 'Controles existentes',
                  'Classificação do risco (O x D x G = NPR)', 'Ações para reduzir a ocorrência', 'Proposta NQSP']
    for campo in campos_npr:
        p = document.add_paragraph()
        p.add_run(f'{campo}: ').bold = True
        p.add_run('[preencher]')

    # --- SEÇÃO 11 — AÇÕES TÁTICAS E ESTRATÉGICAS ---
    _secao(document, f'11. AÇÕES TÁTICAS E ESTRATÉGICAS DESENVOLVIDAS NO {trimestre_atual}º TRIMESTRE DE {ano_atual}')

    document.add_paragraph('Ações concluídas referente ao relatório trimestral anterior:').runs[0].bold = True
    if acoes_taticas.get('acoes_anteriores'):
        tab_ant = document.add_table(rows=1, cols=4)
        tab_ant.style = 'Table Grid'
        for i, col in enumerate(['Nº', 'Ação', 'Status', 'Responsável']):
            tab_ant.rows[0].cells[i].text = col
        for idx, acao in enumerate(acoes_taticas['acoes_anteriores'], start=1):
            row = tab_ant.add_row().cells
            row[0].text = str(idx)
            row[1].text = acao.get('acao', '')
            row[2].text = acao.get('status', '')
            row[3].text = acao.get('responsavel', '')

    document.add_paragraph(f'Ações {trimestre_atual}º Trimestre de {ano_atual}:').runs[0].bold = True
    if acoes_taticas.get('acoes_novas'):
        tab_nov = document.add_table(rows=1, cols=5)
        tab_nov.style = 'Table Grid'
        for i, col in enumerate(['Nº', 'Origem', 'Ação', 'Status', 'Responsável']):
            tab_nov.rows[0].cells[i].text = col
        for idx, acao in enumerate(acoes_taticas['acoes_novas'], start=1):
            row = tab_nov.add_row().cells
            row[0].text = str(idx)
            row[1].text = acao.get('origem', '')
            row[2].text = acao.get('acao', '')
            row[3].text = acao.get('status', '')
            row[4].text = acao.get('responsavel', '')

    # --- SEÇÃO 12 — REFERÊNCIAS ---
    _secao(document, '12. REFERÊNCIAS')
    referencias = [
        'Azevedo et al. Prevalência de incidentes relacionados à medicação em UTI. Acta Paulista de Enfermagem, 2015; 28(4): 331-336.',
        'Shehata et al. Descriptive analysis of medication errors. Journal of the American Medical Informatics Association, 2015; 23(2): 366-374.',
        'Raleigh VS et al. Patient safety indicators for England. BMJ 2008; 337:a1702.',
        'Runciman W et al. Towards an International Classification for Patient Safety. Int J Qual Health Care 2009; 21:18-26.',
        'World Health Organization. World Alliance for Patient Safety. Genebra: 2009.',
    ]
    for ref in referencias:
        document.add_paragraph(ref, style='List Bullet')

    # --- RODAPÉ ---
    document.add_paragraph(f'\n{trimestre_atual}º Trimestre de {ano_atual}')
    document.add_paragraph(f'Elaboração: {hospital} — Núcleo de Qualidade e Segurança do Paciente')

    # --- SALVAR ---
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        document.save(tmp.name)
    return tmp.name